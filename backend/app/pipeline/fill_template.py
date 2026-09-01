"""Fill a USER's own Word document with a meeting's data - no code tags needed.

Unlike the built-in docxtpl template (which uses {{ Jinja }} tags), a user
brings an ordinary .docx and we fill it WITHOUT learning any markup. For each
field we look, in this order:

  1. HEADING / LABEL - a section heading or short label that names the field.
     Matching is forgiving: an exact name ("Action Items", "决策"), a label that
     CONTAINS the name ("Key Actions", "put your action items here"), or a close
     typo ("Actoin Items") all resolve; list/table content is filled beneath it.
  2. TABLE LABEL - a details row like "Date | (blank)": the empty cell beside a
     recognised scalar label is filled in place, so a details block needs no
     markers at all.
  3. MARKER - still supported for precise inline placement: a plain-English
     [[Date]] or [[Summary]] dropped exactly where the value should go.

Recognised names work in English, Chinese and Bahasa Melayu, and are
case-insensitive. Loose matching is gated to short, label-like lines so it never
fires inside a sentence. We only INSERT content (never restyle), so fonts and
layout are preserved.

`fill_user_document` returns the saved path + a report of which fields were
placed (and how) and which had data but found no spot. Pass `combined=True` to
fill the SAME document with a whole follow-up chain instead of one meeting:
every list is pooled across the series and tagged with the meeting it came
from, so an organisation keeps its own layout for combined minutes too.
"""
from __future__ import annotations

import difflib
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from app.config import settings
from app.pipeline.export import (
    TemplateRenderError,
    build_combined_context,
    build_meeting_context,
    output_filename,
)

# --- recognised field names (headings AND [[markers]]) ---
# English + Chinese + Bahasa Melayu (added 2026-08-07, so a Malaysian
# organisation can lay its house template out in its own working language).
_FIELD_SYNONYMS = {
    "title": ["meeting title", "title", "meeting name", "subject", "meeting subject",
              "会议标题", "标题", "会议主题", "主题",
              "tajuk", "tajuk mesyuarat", "nama mesyuarat"],
    "date": ["date", "meeting date", "日期", "会议日期",
             "tarikh", "tarikh mesyuarat"],
    "duration": ["duration", "length", "会议时长", "时长",
                 "tempoh", "tempoh mesyuarat", "jangka masa"],
    "language": ["language", "languages", "语言", "bahasa"],
    "participants": ["participant count", "number of participants", "no of participants",
                     "no. of participants", "headcount", "人数", "参与人数", "与会人数", "出席人数",
                     "bilangan peserta", "jumlah peserta", "bilangan kehadiran"],
    "attendees": ["attendees", "attendee", "participants", "participant", "present",
                  "attendance", "出席", "出席人员", "与会者", "参与者", "参会人员", "参加人员",
                  "kehadiran", "hadirin", "peserta", "yang hadir", "ahli hadir",
                  "senarai kehadiran"],
    "summary": ["summary", "meeting summary", "overview", "abstract", "executive summary",
                "摘要", "会议摘要", "概要", "总结", "会议总结", "内容摘要",
                "ringkasan", "ringkasan mesyuarat", "rumusan", "intisari"],
    "action_items": ["action items", "action item", "actions", "action points", "action plan",
                     "tasks", "task list", "to-do", "to do", "todo", "follow-ups", "follow ups",
                     "行动项", "行动事项", "待办", "待办事项", "任务", "任务清单", "行动计划", "跟进事项",
                     "tindakan", "senarai tindakan", "tindakan susulan", "perkara tindakan",
                     "tugasan", "senarai tugasan"],
    "decisions": ["key decisions", "decisions", "decision", "decisions made", "resolutions",
                  "agreements", "决策", "决定", "关键决策", "议决", "决议", "达成的决定",
                  "keputusan", "keputusan penting", "ketetapan", "resolusi"],
    "deadlines": ["deadlines", "deadline", "due dates", "due date", "timeline", "timelines",
                  "key dates", "schedule", "期限", "截止日期", "截止时间", "时间表", "时间节点", "重要日期",
                  "tarikh akhir", "tarikh tutup", "garis masa", "jadual", "tempoh masa"],
    "issues": ["technical issues", "issues", "issue", "problems", "problem", "roadblocks",
               "blockers", "challenges", "技术问题", "问题", "障碍", "技术难题", "待解决问题", "难点",
               "isu", "isu teknikal", "masalah", "permasalahan", "kekangan"],
    "risks": ["risks", "risk", "risks and mitigations", "risk and mitigation",
              "风险", "风险与缓解", "风险点", "潜在风险", "隐患",
              "risiko", "potensi risiko", "risiko dan langkah"],
    # Progress on the PREVIOUS meeting's action items - filled only when this
    # meeting is linked as a follow-up. "Matters arising" and its Malay
    # equivalent "Perkara berbangkit" are the headings real corporate minutes
    # actually use for this section, so they matter more than a literal
    # translation of "carry forward".
    "carry_forward": ["matters arising", "progress since last meeting",
                      "previous action items", "previous actions", "outstanding actions",
                      "carry forward", "carried forward", "follow-up on previous meeting",
                      "review of previous minutes",
                      "上次会议跟进", "上次会议回顾", "前次行动项", "上次行动事项",
                      "perkara berbangkit", "kemajuan tindakan lepas",
                      "tindakan mesyuarat lepas"],
    # NOTE: the words a document uses to NAME ITSELF are deliberately absent
    # here - "minutes", "会议记录"/"记录", "rekod mesyuarat", "catatan". They are
    # what a minutes template calls its own title, not what it calls a verbatim
    # section, and including them meant a heading like "MINUTES OF MEETING"
    # claimed the transcript field and had the whole transcript poured under the
    # document title (BUGS.md #34). Every unambiguous name is still here, in all
    # three languages, so a real transcript section still fills.
    "transcript": ["transcript", "full transcript", "meeting transcript", "verbatim",
                   "verbatim record", "full record",
                   "逐字稿", "全文", "完整记录", "会议全文",
                   "transkrip", "transkrip penuh", "minit penuh"],
}
_SCALAR = {"title", "date", "duration", "language", "participants", "attendees"}
_BLOCK = {"summary", "action_items", "decisions", "deadlines", "issues", "risks",
          "carry_forward", "transcript"}

_MARKER = re.compile(r"\[\[\s*(.+?)\s*\]\]")


def _norm(s: str | None) -> str:
    """Normalise a heading/marker name for matching: lowercase, trim, drop a
    trailing colon/period, collapse whitespace."""
    s = (s or "").strip().lower().rstrip(":：.。").strip()
    return re.sub(r"\s+", " ", s)


# normalised synonym -> field key
_LOOKUP: dict[str, str] = {}
for _key, _syns in _FIELD_SYNONYMS.items():
    for _syn in _syns:
        _LOOKUP[_norm(_syn)] = _key

# synonyms longest-first, for containment matching ("put your action items here")
_SYN_BY_LEN = sorted(_LOOKUP.items(), key=lambda kv: -len(kv[0]))
_HAS_CJK = re.compile(r"[㐀-䶿一-鿿]")


def _loose_match(n: str, scalar_only: bool = False) -> str | None:
    """Match a normalised, LABEL-LIKE line to a field beyond an exact synonym:
    first by whole-word containment of a known synonym (longest wins, so
    "put your action items here" -> action_items), then by fuzzy similarity for
    typos ("actoin items"). Callers must gate this to short/label text so it
    never fires inside a body sentence."""
    best_key, best_len = None, 0
    for syn, key in _SYN_BY_LEN:
        if scalar_only and key not in _SCALAR:
            continue
        if len(syn) < 4:  # skip tiny synonyms that would match noise
            continue
        if _HAS_CJK.search(syn):
            hit = syn in n
        else:
            hit = re.search(r"(?<![a-z])" + re.escape(syn) + r"(?![a-z])", n) is not None
        if hit and len(syn) > best_len:
            best_key, best_len = key, len(syn)
    if best_key:
        return best_key
    fk, fr = None, 0.0
    for syn, key in _LOOKUP.items():
        if scalar_only and key not in _SCALAR:
            continue
        r = difflib.SequenceMatcher(None, n, syn).ratio()
        if r > fr:
            fr, fk = r, key
    # 0.90 is deliberately conservative: it catches genuine typos ("actoin items",
    # "summry") but still flags heavier deviations like "Riks" -> risks, so the
    # feedback stays useful rather than papering over a wrong heading.
    return fk if fr >= 0.90 else None


def _labelish(text: str, style_name: str = "") -> bool:
    """A heading or short standalone label (never a full body sentence)."""
    n = _norm(text)
    if not n:
        return False
    if style_name.startswith("Heading") or style_name.startswith("Title"):
        return True
    return len(n.split()) <= 6 and text.strip()[-1:] not in ".!?。！？"


def _match_field(text: str, style_name: str = "", scalar_only: bool = False) -> str | None:
    """Field key for a heading/label: exact synonym, else loose (if label-like)."""
    n = _norm(text)
    if not n:
        return None
    k = _LOOKUP.get(n)
    if k:
        return k if (not scalar_only or k in _SCALAR) else None
    if _labelish(text, style_name):
        return _loose_match(n, scalar_only)
    return None


# --- value rendering ------------------------------------------------------

def _series(ctx: dict) -> bool:
    """True when the context spans a follow-up CHAIN rather than one meeting
    (build_combined_context). A combined document carries an extra "Meeting"
    column, or a bracketed prefix on a bullet, so a reader always sees which
    session a row came from."""
    return bool(ctx.get("combined"))


def _tag(item: dict, combined: bool) -> str:
    """A "[Session 1] " prefix for a bullet/inline item in a combined document."""
    return f"[{item['meeting']}] " if combined and item.get("meeting") else ""


def _cols(combined: bool, item: dict) -> list[str]:
    """The optional leading "Meeting" cell of a combined document's table row."""
    return [item.get("meeting") or "-"] if combined else []


def _scalar_value(key: str, ctx: dict) -> str:
    m = ctx["meeting"]
    return {
        "title": m["title"],
        "date": m["date"],
        "duration": m["duration"],
        "language": m["language_name"],
        "participants": str(m["speaker_count"]),
        "attendees": ", ".join(ctx["attendees"]) or "—",
    }.get(key, "")


def _block_plaintext(key: str, ctx: dict) -> str:
    """A one-line text rendering, used when a block marker sits inline in a
    sentence or inside a table cell (where we can't insert a table/bullets)."""
    combined = _series(ctx)
    if key == "summary":
        return ctx["summary"]
    if key == "action_items":
        return "; ".join(_tag(a, combined) + a["description"]
                         + (f" ({a['owner']})" if a.get("owner") else "")
                         for a in ctx["action_items"]) or "None recorded."
    if key == "deadlines":
        return "; ".join(_tag(d, combined) + d["description"]
                         + (f" ({d['date']})" if d.get("date") else "")
                         for d in ctx["deadlines"]) or "None recorded."
    if key in ("decisions", "issues"):
        return "; ".join(_tag(i, combined) + i["description"]
                         for i in ctx[key]) or "None recorded."
    if key == "risks":
        return "; ".join(_tag(r, combined) + r["description"]
                         + (f" (Mitigation: {r['mitigation']})" if r.get("mitigation") else "")
                         for r in ctx["risks"]) or "None recorded."
    if key == "carry_forward":
        return "; ".join(_tag(c, combined) + f"{c['description']} — {c['status']}"
                         for c in ctx["carry_forward"]) or "No previous meeting linked."
    if key == "transcript":
        # A combined document pools no transcript (build_combined_context) -
        # state that rather than resolving the marker to an empty string.
        if combined:
            return ctx.get("transcript_note", "")
        return " / ".join(f"[{s['time']}] {s['speaker']}: {s['text']}" for s in ctx["transcript"])
    return ""


# --- docx insertion helpers ----------------------------------------------

def _insert_paragraph_after(paragraph: Paragraph, text: str = "", italic: bool = False) -> Paragraph:
    """Insert a new paragraph right after `paragraph`; return the new Paragraph."""
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    if text:
        run = para.add_run(text)
        run.italic = italic
    return para


def _insert_table_after(doc, paragraph: Paragraph, headers: list[str], rows: list[list]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))  # created at end of body...
    try:
        table.style = "Light Grid Accent 1"
    except Exception:  # noqa: BLE001 - style may not exist in an arbitrary user doc
        pass
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    paragraph._p.addnext(table._tbl)  # ...then moved to sit right after the anchor


def _delete_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def _insert_field(doc, anchor: Paragraph, key: str, ctx: dict) -> None:
    """Insert a field's content as block(s) right after `anchor`.

    A combined context (a whole follow-up chain) renders identically except
    that every pooled row names the meeting it came from - an extra leading
    column in a table, a "[Session 1]" prefix on a bullet.
    """
    combined = _series(ctx)
    if key in _SCALAR:
        _insert_paragraph_after(anchor, _scalar_value(key, ctx))
        return
    if key == "summary":
        # A series keeps each meeting's summary under its own bold heading
        # line, in chain order, rather than running them together into one
        # paragraph no reader could attribute.
        if combined and ctx.get("summaries"):
            cur = anchor
            for s in ctx["summaries"]:
                cur = _insert_paragraph_after(cur, f"{s['title']}  ({s['date']})")
                if cur.runs:
                    cur.runs[0].bold = True
                cur = _insert_paragraph_after(cur, s["text"])
            return
        _insert_paragraph_after(anchor, ctx["summary"])
        return
    if key == "action_items":
        items = ctx["action_items"]
        if not items:
            _insert_paragraph_after(anchor, "None recorded.", italic=True)
            return
        rows = [[str(i)] + _cols(combined, a) +
                [a["description"] + (f"  [{a['source']}]" if a.get("source") else ""),
                 a.get("owner") or "-", a.get("due") or "-"]
                for i, a in enumerate(items, 1)]
        _insert_table_after(
            doc, anchor,
            ["No."] + (["Meeting"] if combined else []) + ["Action", "Owner", "Due"],
            rows,
        )
        return
    if key == "deadlines":
        items = ctx["deadlines"]
        if not items:
            _insert_paragraph_after(anchor, "None recorded.", italic=True)
            return
        rows = [[str(i)] + _cols(combined, d) +
                [d["description"] + (f"  [{d['source']}]" if d.get("source") else ""),
                 d.get("date") or "-"] for i, d in enumerate(items, 1)]
        _insert_table_after(
            doc, anchor,
            ["No."] + (["Meeting"] if combined else []) + ["Deadline", "Date"],
            rows,
        )
        return
    if key == "carry_forward":
        items = ctx["carry_forward"]
        if not items:
            _insert_paragraph_after(
                anchor, "No previous meeting is linked to this one.", italic=True
            )
            return
        rows = [[str(i)] + _cols(combined, c) +
                [c["description"] + (f"  [{c['source']}]" if c.get("source") else ""),
                 c.get("owner") or "-", c["status"], c.get("note") or "-"]
                for i, c in enumerate(items, 1)]
        _insert_table_after(
            doc, anchor,
            ["No."] + (["Reviewed in"] if combined else [])
            + ["Action agreed previously", "Owner", "Status", "Evidence"],
            rows,
        )
        return
    if key in ("decisions", "issues", "risks"):
        items = ctx[key]
        if not items:
            _insert_paragraph_after(anchor, "None recorded.", italic=True)
            return
        cur = anchor
        for it in items:
            line = "• " + _tag(it, combined) + it["description"]
            if it.get("source"):
                line += f"  [{it['source']}]"
            if key == "risks" and it.get("mitigation"):
                line += "  —  Mitigation: " + it["mitigation"]
            cur = _insert_paragraph_after(cur, line)
        return
    if key == "transcript":
        # A combined document deliberately pools no transcript - leaving the
        # user's Transcript heading bare would read as a failure to fill it.
        if combined:
            _insert_paragraph_after(anchor, ctx.get("transcript_note", ""), italic=True)
            return
        cur = anchor
        for seg in ctx["transcript"]:
            cur = _insert_paragraph_after(cur, f"[{seg['time']}] {seg['speaker']}: {seg['text']}")
        return


def _set_paragraph_text(para: Paragraph, new_text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    if para.text == new_text:
        return
    if para.runs:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(new_text)


def _resolve_inline(text: str, ctx: dict, filled: dict) -> str:
    """Replace every recognised [[marker]] in a line with its value (scalar, or
    a one-line rendering for block fields). Unknown markers are left untouched."""
    def repl(match: re.Match) -> str:
        key = _LOOKUP.get(_norm(match.group(1)))
        if not key:
            return match.group(0)
        filled.setdefault(key, "marker")
        return _block_plaintext(key, ctx) if key in _BLOCK else _scalar_value(key, ctx)
    return _MARKER.sub(repl, text)


# --- the two passes -------------------------------------------------------

def _fill_markers(doc, ctx: dict, filled: dict) -> None:
    # Body paragraphs: a line that is ONLY a block marker becomes the real block.
    for para in list(doc.paragraphs):
        only = _MARKER.fullmatch(para.text.strip())
        if only:
            key = _LOOKUP.get(_norm(only.group(1)))
            if key in _BLOCK:
                if key not in filled:
                    _insert_field(doc, para, key, ctx)
                    filled[key] = "marker"
                _delete_paragraph(para)
                continue
        if "[[" in para.text:
            _set_paragraph_text(para, _resolve_inline(para.text, ctx, filled))
    # Table cells: scalar / inline-block markers only (no block insertion).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if "[[" in para.text:
                        _set_paragraph_text(para, _resolve_inline(para.text, ctx, filled))


def _set_cell_value(cell, text: str) -> None:
    """Write text into a table cell, keeping the cell's existing formatting."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def _fill_table_labels(doc, ctx: dict, filled: dict) -> None:
    """Marker-free scalars: a details row like 'Date | (blank)' has the blank
    cell to its right filled in. Only scalar fields, and only when the target
    cell is empty (or holds a leftover [[marker]]), so real data is never
    overwritten."""
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for i in range(len(cells) - 1):
                key = _match_field(cells[i].text, scalar_only=True)
                if not key or key in filled:
                    continue
                target = cells[i + 1]
                if cells[i]._tc is target._tc:  # merged across columns, not a value cell
                    continue
                stripped = target.text.strip()
                if stripped == "" or _MARKER.fullmatch(stripped):
                    _set_cell_value(target, _scalar_value(key, ctx))
                    filled[key] = "table-label"


def _fill_headings(doc, ctx: dict, filled: dict) -> None:
    for para in list(doc.paragraphs):
        key = _match_field(para.text, para.style.name if para.style else "")
        if key and key not in filled:
            _insert_field(doc, para, key, ctx)
            filled[key] = "heading"


# --- public API -----------------------------------------------------------

def recognised_fields(template_path: Path) -> list[str]:
    """Field keys the document references via a heading or a [[marker]]. Used to
    validate an upload has at least one fillable spot. Raises TemplateRenderError
    if the file cannot be opened as a .docx."""
    try:
        doc = Document(str(template_path))
    except Exception as exc:  # noqa: BLE001 - zipfile/docx raise various types
        raise TemplateRenderError(
            f"Could not open this file as a Word .docx ({type(exc).__name__}). "
            "Please upload a valid .docx saved from Word."
        ) from exc
    found: set[str] = set()
    # body paragraphs: heading/label match (exact, containment or fuzzy) + markers
    for para in doc.paragraphs:
        k = _match_field(para.text, para.style.name if para.style else "")
        if k:
            found.add(k)
        for name in _MARKER.findall(para.text):
            key = _LOOKUP.get(_norm(name))
            if key:
                found.add(key)
    # tables: marker-free label->next-cell scalars + inline markers
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for i, cell in enumerate(cells):
                for name in _MARKER.findall(cell.text):
                    key = _LOOKUP.get(_norm(name))
                    if key:
                        found.add(key)
                if i < len(cells) - 1:
                    k = _match_field(cell.text, scalar_only=True)
                    if k and cells[i]._tc is not cells[i + 1]._tc:
                        found.add(k)
    return sorted(found)


def _data_fields(ctx: dict) -> list[str]:
    """Fields that actually have content worth placing (for the skipped report)."""
    fields = ["title", "date", "duration", "language", "participants", "summary"]
    # A combined document deliberately carries no pooled transcript, so a
    # template without a Transcript heading is not "missing" anything.
    if not _series(ctx):
        fields.append("transcript")
    if ctx["attendees"]:
        fields.append("attendees")
    # carry_forward is listed only when this meeting IS a follow-up - otherwise
    # every ordinary template would be reported as "missing" a section that has
    # no data to put in it.
    for key in ("action_items", "decisions", "deadlines", "issues", "risks", "carry_forward"):
        if ctx[key]:
            fields.append(key)
    return fields


def fill_user_document(
    meeting_id: int,
    template_path: Path,
    *,
    template_name: str | None = None,
    combined: bool = False,
) -> tuple[Path, dict]:
    """Fill the user's .docx with a meeting's data via markers + headings.

    `combined=True` fills it with the WHOLE follow-up chain ending at
    `meeting_id` rather than that meeting alone, so an organisation's own
    layout is kept for combined minutes as well as ordinary ones.

    `template_name` is woven into the output filename. Without it a templated
    export saved under exactly the same name as the built-in one, so a browser
    kept the first download and named the second "... (1).docx" - opening the
    familiar name showed the DEFAULT layout and the chosen template looked as
    though it had been ignored.

    Returns (output_path, report) where report = {placed: {field: 'marker'|'heading'},
    skipped: [fields that had data but found no marker/heading]}.
    """
    # ValueError if the meeting is missing or has no transcript yet.
    ctx = build_combined_context(meeting_id) if combined else build_meeting_context(meeting_id)
    try:
        doc = Document(str(template_path))
    except Exception as exc:  # noqa: BLE001
        raise TemplateRenderError(
            f"Could not open this file as a Word .docx ({type(exc).__name__}). "
            "Please upload a valid .docx saved from Word."
        ) from exc

    filled: dict[str, str] = {}
    _fill_markers(doc, ctx, filled)
    _fill_table_labels(doc, ctx, filled)
    _fill_headings(doc, ctx, filled)

    skipped = [k for k in _data_fields(ctx) if k not in filled]

    out_dir = settings.output_dir / str(meeting_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / output_filename(
        ctx.get("file_title") or ctx["meeting"]["title"],
        combined=combined,
        template_name=template_name,
    )
    doc.save(str(out_path))
    return out_path, {"placed": filled, "skipped": skipped}
