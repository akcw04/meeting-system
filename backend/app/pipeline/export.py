"""Word document export (Phase 8).

A meeting becomes a .docx three ways:
  - the built-in default template (templates/default.docx, docxtpl {{ tags }}),
    rendered by export_docx();
  - a USER's own document, filled WITHOUT any code-like syntax by
    pipeline/fill_template.py (plain section headings or friendly [[markers]]);
  - a COMBINED minutes document spanning a follow-up chain of meetings,
    built by export_combined_docx() (Session 22).

Data shaping is shared: build_meeting_context() for one meeting,
build_combined_context() for a whole follow-up chain. Either can feed either
renderer - this module owns the data, the template owns the formatting (the
IR's Separation-of-Concerns design). So a combined export honours the user's
own document too: export_combined_docx() lays out the built-in house style,
while fill_user_document(..., combined=True) pours the same pooled data into
the user's layout.

Output: data/outputs/<meeting_id>/<slug>[_combined]_minutes[_<template>].docx
Re-rendered on every export request so transcript/insight edits are reflected.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docxtpl import DocxTemplate

from app.config import settings
from app.db import get_conn

TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / "default.docx"

LANGUAGE_NAMES = {"en": "English", "zh": "Chinese", "ms": "Malay"}


class TemplateRenderError(ValueError):
    """A Word document could not be produced/filled with a meeting's data.

    Raised for the built-in template (bad Jinja) or a user document we cannot
    open as a .docx. Subclasses ValueError so routes can catch it specifically
    and return a user-fixable 400 instead of a 409.
    """


def _fmt_time(seconds: float | None) -> str:
    s = int(seconds or 0)
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"


def _slug(text: str, max_len: int = 40) -> str:
    text = re.sub(r"[^\w\s-]", "", text).strip().lower()
    text = re.sub(r"[\s_-]+", "_", text)
    return text[:max_len] or "meeting"


def output_filename(
    title: str, *, combined: bool = False, template_name: str | None = None
) -> str:
    """Filename for an exported document.

    The chosen template's name is part of it so a templated export and the
    built-in one never collide. They used to share a name, so a browser saved
    the second download as "... (1).docx" and opening the familiar name showed
    the DEFAULT layout - making the system look as though it had ignored the
    template the user picked.
    """
    suffix = ""
    if template_name:
        suffix = _slug(template_name, 24)
        # A truncated template name shouldn't end mid-word ("..._templa"), so
        # drop the partial trailing token when the slug was actually cut.
        if "_" in suffix and len(_slug(template_name, 200)) > len(suffix):
            suffix = suffix.rsplit("_", 1)[0]
    parts = [_slug(title), "combined" if combined else "", "minutes", suffix]
    return "_".join(p for p in parts if p) + ".docx"



def local_date(stored: str | None) -> str:
    """The calendar date of a stored timestamp, in the machine's own timezone.

    SQLite's CURRENT_TIMESTAMP is UTC and is stored WITHOUT a zone marker
    ("2026-09-19 10:20:46"). Slicing the first ten characters therefore yields
    the UTC date, which is a day early for any meeting uploaded before the
    UTC offset (before 08:00 in Malaysia, UTC+8) - putting the wrong date on
    minutes that get circulated. Convert first, then take the date.

    Falls back to the raw first ten characters if the value is missing or in an
    unexpected shape, so an export never fails over a timestamp.
    """
    if not stored:
        return ""
    text = str(stored).strip().replace("T", " ").removesuffix("Z")
    try:
        naive = datetime.strptime(text[:19], "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return str(stored)[:10]
    return naive.replace(tzinfo=timezone.utc).astimezone().strftime("%Y-%m-%d")


def build_meeting_context(meeting_id: int) -> dict:
    """Shape a meeting's DB rows into the render context shared by the built-in
    docxtpl template AND the user-document fill engine.

    Raises ValueError if the meeting is missing or has no transcript yet.
    """
    with get_conn() as conn:
        meeting = conn.execute(
            "SELECT * FROM meetings WHERE id = ?", (meeting_id,)
        ).fetchone()
        if meeting is None:
            raise ValueError(f"Meeting {meeting_id} not found")

        speakers = conn.execute(
            "SELECT * FROM speakers WHERE meeting_id = ? ORDER BY id", (meeting_id,)
        ).fetchall()
        segments = conn.execute(
            "SELECT * FROM segments WHERE meeting_id = ? ORDER BY start_seconds",
            (meeting_id,),
        ).fetchall()

        def rows(table: str) -> list[dict]:
            return [
                dict(r)
                for r in conn.execute(
                    f"SELECT * FROM {table} WHERE meeting_id = ? ORDER BY id",
                    (meeting_id,),
                ).fetchall()
            ]

        action_items = rows("action_items")
        decisions = rows("decisions")
        deadlines = rows("deadlines")
        issues = rows("issues")
        risks = rows("risks")

    if not segments:
        raise ValueError(
            f"Meeting {meeting_id} has no transcript yet (status: {meeting['status']})"
        )

    label_of = {s["id"]: (s["display_name"] or s["label"]) for s in speakers}
    time_of = {s["id"]: _fmt_time(s["start_seconds"]) for s in segments}  # for source refs

    def _src(row: dict) -> str | None:
        return time_of.get(row.get("source_segment_id"))

    # Carry-forward rows exist only when this meeting was linked as the
    # follow-up of an earlier one; otherwise this is simply an empty list and
    # every template renders exactly as it did before the feature existed.
    # Local import: carryforward -> categorize -> config, so importing it at
    # module scope would make this module's import order matter.
    from app.pipeline.carryforward import load_carry_forward

    carry_forward = [
        {
            "description": r["description"],
            "owner": r["owner"],
            "status": CARRY_FORWARD_LABELS.get(r["status"], r["status"]),
            "note": r["note"],
            "source": _src(r),
        }
        for r in load_carry_forward(meeting_id)
    ]

    return {
        "meeting": {
            "title": meeting["title"],
            "date": local_date(meeting["created_at"]),
            "duration": _fmt_time(meeting["duration_seconds"]),
            "language_name": LANGUAGE_NAMES.get(meeting["language"], meeting["language"] or "-"),
            "speaker_count": len(speakers),
        },
        "attendees": [label_of[s["id"]] for s in speakers],
        "summary": meeting["summary"] or "No summary generated yet.",
        "action_items": [
            {"description": a["description"], "owner": a["owner"], "due": a["due_date"], "source": _src(a)}
            for a in action_items
        ],
        "decisions": [{"description": d["description"], "source": _src(d)} for d in decisions],
        "deadlines": [
            {"description": d["description"], "date": d["target_date"], "source": _src(d)} for d in deadlines
        ],
        "issues": [{"description": i["description"], "source": _src(i)} for i in issues],
        "risks": [
            {"description": r["description"], "mitigation": r["mitigation"], "source": _src(r)} for r in risks
        ],
        "carry_forward": carry_forward,
        "transcript": [
            {
                "time": _fmt_time(seg["start_seconds"]),
                "speaker": label_of.get(seg["speaker_id"], "Unknown"),
                "text": seg["text"].strip(),
            }
            for seg in segments
        ],
        "generated_on": datetime.now().strftime("%d %B %Y, %H:%M"),
    }


def export_docx(meeting_id: int, template_path: Path | None = None) -> Path:
    """Render the meeting minutes via the built-in docxtpl template and return
    the output path. (User-uploaded documents go through fill_template.py.)
    """
    template = Path(template_path) if template_path else TEMPLATE_PATH
    if not template.is_file():
        raise FileNotFoundError(
            f"Template not found at {template}. Regenerate it with: "
            f"python templates\\generate_default_template.py"
        )

    context = build_meeting_context(meeting_id)
    doc = DocxTemplate(str(template))
    try:
        doc.render(context)
    except Exception as exc:  # malformed Jinja in the built-in template
        raise TemplateRenderError(
            f"Could not render the built-in template ({type(exc).__name__}: {exc})."
        ) from exc

    out_dir = settings.output_dir / str(meeting_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / output_filename(context["meeting"]["title"])
    doc.save(str(out_path))
    return out_path


# === Combined minutes across a follow-up chain (Session 22) ===

_ACCENT = RGBColor(0x1F, 0x4E, 0x79)  # same corporate blue as the built-in template

# How each carry-forward verdict is worded in the document. The DB stores a
# stable machine value; only this map decides what the reader sees.
CARRY_FORWARD_LABELS = {
    "completed": "Completed",
    "in_progress": "In progress",
    "blocked": "Blocked",
    "changed": "Changed",
    "not_discussed": "Not discussed",
}


def _style_base(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for level, size in (("Heading 1", 16), ("Heading 2", 13)):
        st = doc.styles[level]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = _ACCENT
        st.font.bold = True


def _table(doc, headers: list[str], rows: list[list[str]]):
    """A styled table with a bold header row. Returns None when there is
    nothing to show, so callers can write a placeholder line instead."""
    if not rows:
        return None
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:  # style missing from a stripped-down python-docx default
        pass
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if value is None else str(value)
    return table


def _none_recorded(doc) -> None:
    doc.add_paragraph("None recorded.").runs[0].italic = True


def _unique(values):
    """Order-preserving de-duplication - keeps a series' attendee/language
    lists in the order they were first seen rather than alphabetised."""
    seen, out = set(), []
    for v in values:
        if v and v not in seen:
            seen.add(v)
            out.append(v)
    return out


def build_combined_context(meeting_id: int) -> dict:
    """Shape a whole follow-up CHAIN into the same context shape as
    build_meeting_context, so a user's own document can be filled with a series
    of meetings as well as with a single one.

    Everything is pooled in chain order and tagged with the meeting it came
    from, so no outcome is silently attributed to the wrong session:

      - the scalars describe the SERIES - a date range, the total duration,
        every language and every attendee seen across it;
      - action items, decisions, deadlines, issues, risks and the carry-forward
        verdicts are concatenated, each carrying a 'meeting' key the fill engine
        renders as its own column (tables) or prefix (bullets);
      - 'summaries' holds one entry per meeting, in order; 'summary' is the same
        content flattened for a marker sitting inline in a sentence.

    Transcripts are deliberately left OUT, exactly as the built-in combined
    layout leaves them out: each meeting's own export already carries its
    verbatim record, and pooling several hours of speech would swamp the
    minutes. 'transcript' is therefore empty and 'transcript_note' states the
    omission in the document itself rather than leaving a silent gap.

    Raises ValueError if the meeting is missing or has no transcript yet.
    """
    # Local import: carryforward -> categorize -> config, so importing it at
    # module scope would make this module's import order matter.
    from app.pipeline.carryforward import meeting_chain

    chain = meeting_chain(meeting_id)  # oldest first
    contexts = {mid: build_meeting_context(mid) for mid in chain}
    latest = contexts[meeting_id]["meeting"]
    n = len(chain)

    with get_conn() as conn:
        durations = conn.execute(
            "SELECT id, duration_seconds FROM meetings "
            f"WHERE id IN ({','.join('?' * n)})",
            chain,
        ).fetchall()
    total_seconds = sum((r["duration_seconds"] or 0) for r in durations)

    dates = [contexts[mid]["meeting"]["date"] for mid in chain]
    span = dates[0] if dates[0] == dates[-1] else f"{dates[0]} to {dates[-1]}"

    def pooled(key: str) -> list[dict]:
        """Every item of one category across the series, each tagged with the
        meeting it came from."""
        return [
            {**item, "meeting": contexts[mid]["meeting"]["title"]}
            for mid in chain
            for item in contexts[mid][key]
        ]

    summaries = [
        {
            "title": contexts[mid]["meeting"]["title"],
            "date": contexts[mid]["meeting"]["date"],
            "text": contexts[mid]["summary"],
        }
        for mid in chain
    ]
    attendees = _unique(a for mid in chain for a in contexts[mid]["attendees"])

    return {
        # The fill engine branches on this: a combined document needs the
        # meeting each row came from, a single-meeting one does not.
        "combined": True,
        "chain_length": n,
        # The display title below is deliberately wordy; the filename uses the
        # plain meeting title so downloads stay readable.
        "file_title": latest["title"],
        "meeting": {
            # Named as a series so a reader never mistakes a pooled document
            # for the minutes of the latest meeting alone.
            "title": latest["title"] if n == 1
                     else f"{latest['title']} (combined series of {n} meetings)",
            "date": span,
            "duration": _fmt_time(total_seconds),
            "language_name": ", ".join(
                _unique(contexts[mid]["meeting"]["language_name"] for mid in chain)
            ) or "-",
            "speaker_count": len(attendees),
        },
        "attendees": attendees,
        "summary": "\n".join(
            f"{s['title']} ({s['date']}): {s['text']}" for s in summaries
        ),
        "summaries": summaries,
        "action_items": pooled("action_items"),
        "decisions": pooled("decisions"),
        "deadlines": pooled("deadlines"),
        "issues": pooled("issues"),
        "risks": pooled("risks"),
        "carry_forward": pooled("carry_forward"),
        "transcript": [],
        "transcript_note": (
            "Full transcripts are not reproduced in a combined document - "
            f"each of the {n} meeting{'s' if n != 1 else ''} in this series "
            "carries its verbatim record in its own export."
        ),
        "generated_on": datetime.now().strftime("%d %B %Y, %H:%M"),
    }


def export_combined_docx(meeting_id: int) -> Path:
    """Produce ONE minutes document covering a meeting and everything it follows.

    Walks the follow-up chain back to its first meeting, then writes:

      1. a series overview - every meeting in order, with date and duration;
      2. progress on previous actions - the carry-forward verdicts for each
         follow-up step, so a reader sees at a glance what was actually
         delivered between meetings;
      3. consolidated outcomes - action items, decisions, deadlines, issues and
         risks pooled across the whole series, each tagged with the meeting it
         came from, so nothing is silently attributed to the wrong session;
      4. each meeting's own summary, in order.

    The full transcripts are deliberately NOT included: this is a minutes
    document spanning several meetings, and each meeting's individual export
    already carries its verbatim record.

    Raises ValueError when the meeting is missing or has no transcript.
    """
    # Imported here rather than at module scope: carryforward imports from
    # categorize, which imports settings - keeping it local avoids a circular
    # import at startup while export stays importable on its own.
    from app.pipeline.carryforward import load_carry_forward, meeting_chain

    chain = meeting_chain(meeting_id)
    contexts = {mid: build_meeting_context(mid) for mid in chain}
    latest = contexts[meeting_id]

    doc = Document()
    _style_base(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Combined Minutes of Meeting")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = _ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(latest["meeting"]["title"]).font.size = Pt(14)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    n = len(chain)
    caption.add_run(
        f"{n} meeting{'s' if n != 1 else ''} in this series"
        f"  ·  generated {datetime.now().strftime('%d %B %Y, %H:%M')}"
    ).italic = True

    # --- 1. series overview -------------------------------------------------
    doc.add_heading("Meeting Series", level=1)
    _table(
        doc,
        ["No.", "Meeting", "Date", "Duration", "Participants"],
        [
            [
                str(i),
                contexts[mid]["meeting"]["title"],
                contexts[mid]["meeting"]["date"],
                contexts[mid]["meeting"]["duration"],
                str(contexts[mid]["meeting"]["speaker_count"]),
            ]
            for i, mid in enumerate(chain, start=1)
        ],
    )

    # --- 2. progress on previous actions ------------------------------------
    doc.add_heading("Progress on Previous Actions", level=1)
    any_progress = False
    for position, mid in enumerate(chain):
        rows = load_carry_forward(mid)
        if not rows:
            continue
        any_progress = True
        previous_title = (
            contexts[chain[position - 1]]["meeting"]["title"] if position else "the previous meeting"
        )
        doc.add_heading(
            f"Reviewed in: {contexts[mid]['meeting']['title']}"
            f"  (carried from: {previous_title})",
            level=2,
        )
        _table(
            doc,
            ["No.", "Action agreed previously", "Owner", "Status", "Evidence"],
            [
                [
                    str(i),
                    r["description"],
                    r["owner"] or "-",
                    CARRY_FORWARD_LABELS.get(r["status"], r["status"]),
                    r["note"] or "-",
                ]
                for i, r in enumerate(rows, start=1)
            ],
        )
    if not any_progress:
        doc.add_paragraph(
            "No follow-up analysis has been run for this series yet."
        ).runs[0].italic = True

    # --- 3. consolidated outcomes ------------------------------------------
    doc.add_heading("Consolidated Outcomes", level=1)

    def pooled(key: str):
        """Every item of one category across the series, tagged with its meeting."""
        for mid in chain:
            for item in contexts[mid][key]:
                yield contexts[mid]["meeting"]["title"], item

    doc.add_heading("Action Items", level=2)
    rows = [
        [str(i), title_, it["description"], it.get("owner") or "-", it.get("due") or "-"]
        for i, (title_, it) in enumerate(pooled("action_items"), start=1)
    ]
    if not _table(doc, ["No.", "Meeting", "Action", "Owner", "Due"], rows):
        _none_recorded(doc)

    doc.add_heading("Key Decisions", level=2)
    rows = [
        [str(i), title_, it["description"]]
        for i, (title_, it) in enumerate(pooled("decisions"), start=1)
    ]
    if not _table(doc, ["No.", "Meeting", "Decision"], rows):
        _none_recorded(doc)

    doc.add_heading("Deadlines", level=2)
    rows = [
        [str(i), title_, it["description"], it.get("date") or "-"]
        for i, (title_, it) in enumerate(pooled("deadlines"), start=1)
    ]
    if not _table(doc, ["No.", "Meeting", "Deadline", "Date"], rows):
        _none_recorded(doc)

    doc.add_heading("Technical Issues", level=2)
    rows = [
        [str(i), title_, it["description"]]
        for i, (title_, it) in enumerate(pooled("issues"), start=1)
    ]
    if not _table(doc, ["No.", "Meeting", "Issue"], rows):
        _none_recorded(doc)

    doc.add_heading("Risks", level=2)
    rows = [
        [str(i), title_, it["description"], it.get("mitigation") or "-"]
        for i, (title_, it) in enumerate(pooled("risks"), start=1)
    ]
    if not _table(doc, ["No.", "Meeting", "Risk", "Mitigation"], rows):
        _none_recorded(doc)

    # --- 4. per-meeting summaries -------------------------------------------
    doc.add_heading("Meeting Summaries", level=1)
    for i, mid in enumerate(chain, start=1):
        ctx = contexts[mid]
        doc.add_heading(f"{i}. {ctx['meeting']['title']}  ({ctx['meeting']['date']})", level=2)
        doc.add_paragraph(ctx["summary"])
        doc.add_paragraph("Attendees: " + (", ".join(ctx["attendees"]) or "—"))

    out_dir = settings.output_dir / str(meeting_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / output_filename(latest["meeting"]["title"], combined=True)
    doc.save(str(out_path))
    return out_path
