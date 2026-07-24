"""Generate templates/default.docx - the default meeting-minutes template.

Run me whenever you want to regenerate the default template from scratch:

    python templates\\generate_default_template.py

The output is a normal Word document containing Jinja2 tags that
app/pipeline/export.py fills via docxtpl. Anyone can open default.docx
in Word and restyle it (fonts, colors, logo, ordering) WITHOUT touching
Python - that's the Separation-of-Concerns design from the IR. Keep the
{{ tags }} and {%p/%tr loop markers %} intact when restyling.

Context variables the template expects (provided by export.py):
    meeting: {title, date, duration, language_name, speaker_count}
    attendees: ["Speaker 1", "Mr. Tan", ...]
    summary: str
    action_items: [{description, owner, due}]
    decisions: [{description}]
    deadlines: [{description, date}]
    issues: [{description}]
    risks: [{description, mitigation}]
    transcript: [{time, speaker, text}]
    generated_on: str
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)  # corporate dark blue
OUT = Path(__file__).resolve().parent / "default.docx"


def style_base(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for level, size in (("Heading 1", 16), ("Heading 2", 13)):
        st = doc.styles[level]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.font.bold = True


def add_heading(doc: Document, text: str, level: int = 2):
    return doc.add_heading(text, level=level)


def add_items_table(doc: Document, loop_var: str, columns: list[tuple[str, str]]):
    """Header row + docxtpl row-loop over `loop_var`.

    columns: [(header_text, jinja_expression), ...]
    """
    table = doc.add_table(rows=3, cols=len(columns))
    table.style = "Light Grid Accent 1"
    for i, (header, _) in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    # docxtpl row-loop opener: must live alone in the first cell of its row
    table.rows[1].cells[0].text = "{%tr for item in " + loop_var + " %}"
    for i, (_, expr) in enumerate(columns):
        table.rows[2].cells[i].text = expr
    row_end = table.add_row()
    row_end.cells[0].text = "{%tr endfor %}"
    return table


def add_empty_note(doc: Document, loop_var: str):
    # docxtpl rule: every {%p %} tag must occupy its own paragraph -
    # cramming if/endif into one paragraph breaks Jinja pairing.
    doc.add_paragraph("{%p if not " + loop_var + " %}")
    doc.add_paragraph("None recorded.")
    doc.add_paragraph("{%p endif %}")


def main() -> None:
    doc = Document()
    style_base(doc)

    # === Cover header ===
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Meeting Minutes")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("{{ meeting.title }}").font.size = Pt(14)

    # === Meta table ===
    meta = doc.add_table(rows=5, cols=2)
    meta.style = "Light List Accent 1"
    rows = [
        ("Date", "{{ meeting.date }}"),
        ("Duration", "{{ meeting.duration }}"),
        ("Language", "{{ meeting.language_name }}"),
        ("Participants detected", "{{ meeting.speaker_count }}"),
        ("Attendees", "{{ attendees | join(', ') }}"),
    ]
    for i, (k, v) in enumerate(rows):
        meta.rows[i].cells[0].text = k
        for r in meta.rows[i].cells[0].paragraphs[0].runs:
            r.bold = True
        meta.rows[i].cells[1].text = v

    # === Summary ===
    add_heading(doc, "Meeting Summary", level=1)
    doc.add_paragraph("{{ summary }}")

    # === Action Items ===
    add_heading(doc, "Action Items", level=1)
    add_items_table(
        doc,
        "action_items",
        [
            ("No.", "{{ loop.index }}"),
            ("Action", "{{ item.description }}{% if item.source %}  [{{ item.source }}]{% endif %}"),
            ("Owner", "{{ item.owner or '-' }}"),
            ("Due", "{{ item.due or '-' }}"),
        ],
    )
    add_empty_note(doc, "action_items")

    # === Key Decisions ===
    add_heading(doc, "Key Decisions", level=1)
    doc.add_paragraph("{%p for item in decisions %}")
    doc.add_paragraph("{{ item.description }}{% if item.source %}  [{{ item.source }}]{% endif %}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    add_empty_note(doc, "decisions")

    # === Deadlines ===
    add_heading(doc, "Deadlines", level=1)
    add_items_table(
        doc,
        "deadlines",
        [
            ("No.", "{{ loop.index }}"),
            ("Deadline", "{{ item.description }}{% if item.source %}  [{{ item.source }}]{% endif %}"),
            ("Date", "{{ item.date or '-' }}"),
        ],
    )
    add_empty_note(doc, "deadlines")

    # === Technical Issues ===
    add_heading(doc, "Technical Issues", level=1)
    doc.add_paragraph("{%p for item in issues %}")
    doc.add_paragraph("{{ item.description }}{% if item.source %}  [{{ item.source }}]{% endif %}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    add_empty_note(doc, "issues")

    # === Risks ===
    add_heading(doc, "Risks", level=1)
    doc.add_paragraph("{%p for item in risks %}")
    doc.add_paragraph(
        "{{ item.description }}{% if item.source %}  [{{ item.source }}]{% endif %}"
        "{% if item.mitigation %}  -  Mitigation: {{ item.mitigation }}{% endif %}",
        style="List Bullet",
    )
    doc.add_paragraph("{%p endfor %}")
    add_empty_note(doc, "risks")

    # === Full transcript ===
    doc.add_page_break()
    add_heading(doc, "Full Transcript", level=1)
    doc.add_paragraph("{%p for seg in transcript %}")
    doc.add_paragraph("[{{ seg.time }}] {{ seg.speaker }}: {{ seg.text }}")
    doc.add_paragraph("{%p endfor %}")

    # === Footer line ===
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(
        "Generated by Kairos on {{ generated_on }} - reviewed and approved by the meeting owner."
    )
    fr.font.size = Pt(8)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
