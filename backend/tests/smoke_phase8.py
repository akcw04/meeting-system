"""Phase 8 smoke test - Word document export.

Self-contained: inserts a synthetic fully-processed meeting straight into
the DB (no audio/LLM needed), renders the docx via the export pipeline +
the HTTP route, then reads the document back with python-docx and checks
the planted content actually appears.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
os.chdir(BACKEND)
sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402

from app.db import get_conn, init_db  # noqa: E402
from app.pipeline.export import export_docx  # noqa: E402


def plant_meeting() -> int:
    init_db()
    with get_conn() as conn:
        cur = conn.execute(
            "INSERT INTO meetings (title, original_filename, audio_path, duration_seconds, "
            "language, primary_language, status, summary) "
            "VALUES (?, ?, ?, ?, ?, ?, 'ready', ?)",
            (
                "Phase 8 Export Test Meeting",
                "fake.mp4",
                "data/uploads/none/audio.wav",
                754.0,
                "en",
                "en",
                "The team reviewed sprint progress, agreed to adopt PostgreSQL, "
                "and assigned payment-gateway debugging ahead of the June code freeze.",
            ),
        )
        mid = cur.lastrowid

        spk = {}
        for label, name in (("Speaker 1", "Alice"), ("Speaker 2", None)):
            c = conn.execute(
                "INSERT INTO speakers (meeting_id, label, display_name) VALUES (?, ?, ?)",
                (mid, label, name),
            )
            spk[label] = c.lastrowid

        seg_rows = [
            (0.0, 4.2, "Good morning, let's begin the export test.", spk["Speaker 1"]),
            (4.5, 9.8, "We agreed PostgreSQL is the database going forward.", spk["Speaker 2"]),
            (10.0, 15.0, "Debug the payment gateway by Wednesday please.", spk["Speaker 1"]),
        ]
        seg_ids = []
        for start, end, text, sid in seg_rows:
            c = conn.execute(
                "INSERT INTO segments (meeting_id, speaker_id, start_seconds, end_seconds, text, language) "
                "VALUES (?, ?, ?, ?, ?, 'en')",
                (mid, sid, start, end, text),
            )
            seg_ids.append(c.lastrowid)

        conn.execute(
            "INSERT INTO action_items (meeting_id, description, owner, due_date, source_segment_id) "
            "VALUES (?, 'Debug payment gateway', 'Alice', 'Wednesday', ?)",
            (mid, seg_ids[2]),
        )
        conn.execute(
            "INSERT INTO decisions (meeting_id, description, source_segment_id) "
            "VALUES (?, 'Adopt PostgreSQL as the project database', ?)",
            (mid, seg_ids[1]),
        )
        conn.execute(
            "INSERT INTO deadlines (meeting_id, description, target_date, source_segment_id) "
            "VALUES (?, 'Payment gateway fix', 'Wednesday', ?)",
            (mid, seg_ids[2]),
        )
        conn.execute(
            "INSERT INTO risks (meeting_id, description, mitigation, source_segment_id) "
            "VALUES (?, 'Vendor API may slip', 'Prepare fallback vendor', ?)",
            (mid, seg_ids[1]),
        )
        # issues left empty on purpose - tests the 'None recorded.' path
    return mid


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main() -> int:
    print("=== Phase 8 smoke test: docx export ===")
    mid = plant_meeting()
    print(f"Planted synthetic meeting id={mid}")

    # --- Direct pipeline call ---
    out = export_docx(mid)
    print(f"Rendered: {out} ({out.stat().st_size} bytes)")
    assert out.is_file() and out.stat().st_size > 10_000

    text = docx_text(out)
    expected = [
        "Phase 8 Export Test Meeting",          # title
        "12:34",                                 # duration 754s -> 12:34
        "Alice",                                 # renamed speaker + owner
        "Debug payment gateway",                 # action item
        "Adopt PostgreSQL as the project database",  # decision
        "Vendor API may slip",                   # risk
        "Prepare fallback vendor",               # mitigation
        "None recorded.",                        # empty issues section note
        "Good morning, let's begin the export test.",  # transcript line
    ]
    for needle in expected:
        assert needle in text, f"Missing from docx: {needle!r}"
    # Jinja leftovers would mean a broken template
    for marker in ("{{", "{%"):
        assert marker not in text, f"Unrendered template marker {marker!r} found!"
    print("Direct export content checks PASSED")

    # --- Via the HTTP route ---
    from fastapi.testclient import TestClient
    from app.main import app

    with TestClient(app) as client:
        r = client.get(f"/meetings/{mid}/export/docx")
        assert r.status_code == 200, r.text
        ctype = r.headers["content-type"]
        assert "officedocument.wordprocessingml" in ctype, ctype
        assert len(r.content) > 10_000
        print(f"Route export PASSED ({len(r.content)} bytes, filename in headers: "
              f"{r.headers.get('content-disposition')})")

    print("\nPhase 8 smoke test PASSED.")
    print(f"Open it yourself: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
