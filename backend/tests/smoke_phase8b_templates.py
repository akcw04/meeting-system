"""Phase 8b smoke test - user-document export templates (IR §2.2.7), no-code fill.

Self-contained: plants a synthetic 'ready' meeting, builds a USER doc that uses
BOTH a [[marker]] and plain section HEADINGS, then exercises the feature via the
HTTP routes (TestClient):
  POST   /templates                     upload + "has a fillable spot" validation
  GET    /templates                     listing
  GET    /templates/starter/download    the friendly sample
  GET    /meetings/{id}/export/docx?template_id=   fill the user's doc (marker + headings)
  DELETE /templates/{id}                cleanup

Asserts the meeting's data lands at the marker / under the right headings, that a
doc with no recognised sections is rejected, and cleans up after itself.
"""
from __future__ import annotations

import io
import os
import shutil
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
os.chdir(BACKEND)
sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402

from app.db import get_conn, init_db  # noqa: E402

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def plant_meeting() -> int:
    init_db()
    with get_conn() as conn:
        cur = conn.execute(
            "INSERT INTO meetings (title, original_filename, audio_path, duration_seconds, "
            "language, primary_language, status, summary) "
            "VALUES (?, 'fake.mp4', 'data/uploads/none/audio.wav', 600.0, 'en', 'en', 'ready', ?)",
            ("Phase 8b Template Meeting",
             "A synthetic summary used to validate the no-code template fill."),
        )
        mid = cur.lastrowid
        c = conn.execute(
            "INSERT INTO speakers (meeting_id, label, display_name) VALUES (?, 'Speaker 1', 'Alice')",
            (mid,),
        )
        c = conn.execute(
            "INSERT INTO segments (meeting_id, speaker_id, start_seconds, end_seconds, text, language) "
            "VALUES (?, ?, 0.0, 5.0, 'This is a planted transcript line.', 'en')",
            (mid, c.lastrowid),
        )
        conn.execute(
            "INSERT INTO action_items (meeting_id, description, owner, due_date, source_segment_id) "
            "VALUES (?, 'Planted action item', 'Alice', 'Friday', ?)",
            (mid, c.lastrowid),
        )
    return mid


def cleanup_meeting(mid: int) -> None:
    with get_conn() as conn:
        conn.execute("DELETE FROM meetings WHERE id = ?", (mid,))
    for sub in ("outputs", "uploads"):
        d = BACKEND / "data" / sub / str(mid)
        if d.exists():
            shutil.rmtree(d, ignore_errors=True)


def build_user_doc(path: Path) -> None:
    """A user doc using an inline marker + a standalone block marker + headings."""
    doc = Document()
    doc.add_paragraph("Meeting: [[Meeting Title]]")   # inline scalar marker
    doc.add_paragraph("[[Action Items]]")             # standalone block marker -> table
    doc.add_heading("Summary", level=1)               # heading -> summary fills under it
    doc.add_heading("Attendees", level=1)             # heading -> attendees fill under it
    doc.add_heading("Notes", level=1)                 # NOT recognised -> left alone
    doc.save(str(path))


def docx_text(blob: bytes) -> str:
    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main() -> int:
    print("=== Phase 8b smoke test: no-code user templates ===")
    mid = plant_meeting()
    print(f"Planted synthetic meeting id={mid}")

    fx = BACKEND / "tests" / "fixtures"
    fx.mkdir(parents=True, exist_ok=True)
    good = fx / "_user_template.docx"
    build_user_doc(good)
    empty = fx / "_no_sections.docx"
    d = Document()
    d.add_paragraph("Just some text with no recognised sections.")
    d.save(str(empty))

    from fastapi.testclient import TestClient
    from app.main import app

    tid = None
    try:
        with TestClient(app) as client:
            # 1. Upload the heading/marker doc -> 201.
            r = client.post("/templates",
                            files={"file": ("mine.docx", good.read_bytes(), DOCX_CT)},
                            data={"name": "My Layout"})
            assert r.status_code == 201, r.text
            tid = r.json()["id"]
            print(f"upload (headings+marker)   OK (id={tid})")

            # 2. Reject a doc with no recognised sections, and a non-.docx.
            r = client.post("/templates",
                            files={"file": ("empty.docx", empty.read_bytes(), DOCX_CT)},
                            data={"name": "Empty"})
            assert r.status_code == 400, r.text
            r = client.post("/templates",
                            files={"file": ("bogus.docx", b"not a real docx", "application/octet-stream")},
                            data={"name": "Bogus"})
            assert r.status_code == 400, r.text
            print("reject no-sections + bogus OK (-> 400)")

            # 3. Sample ("starter") download is a valid .docx (zip magic 'PK').
            r = client.get("/templates/starter/download")
            assert r.status_code == 200 and r.content[:2] == b"PK", r.text
            print("sample download            OK")

            # 4. Export the meeting WITH the user doc; data lands at marker + headings.
            r = client.get(f"/meetings/{mid}/export/docx", params={"template_id": tid})
            assert r.status_code == 200, r.text
            txt = docx_text(r.content)
            assert "Phase 8b Template Meeting" in txt, "title marker not filled"
            assert "Planted action item" in txt, "action items (block marker) not filled"
            assert "synthetic summary used to validate" in txt, "summary heading not filled"
            assert "Alice" in txt, "attendees heading not filled"
            assert "[[" not in txt, "a marker was left unfilled"
            print("fill via marker+headings   OK (title, actions, summary, attendees placed)")

            # 5. Default export (no template) still works.
            r = client.get(f"/meetings/{mid}/export/docx")
            assert r.status_code == 200 and r.content[:2] == b"PK"
            print("default export             OK")

            # 6. Bad template id -> 404; delete -> 204; gone from listing + disk.
            assert client.get(f"/meetings/{mid}/export/docx",
                              params={"template_id": 999999}).status_code == 404
            assert client.delete(f"/templates/{tid}").status_code == 204
            assert all(t["id"] != tid for t in client.get("/templates").json())
            assert not (BACKEND / "data" / "templates" / f"{tid}.docx").is_file()
            print("bad id 404 + delete 204    OK")
            tid = None
    finally:
        if tid is not None:
            with get_conn() as conn:
                conn.execute("DELETE FROM templates WHERE id = ?", (tid,))
            (BACKEND / "data" / "templates" / f"{tid}.docx").unlink(missing_ok=True)
        good.unlink(missing_ok=True)
        empty.unlink(missing_ok=True)
        cleanup_meeting(mid)
        print(f"cleaned up planted meeting id={mid}")

    print("\nPhase 8b smoke test PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
