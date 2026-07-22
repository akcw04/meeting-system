"""Stress-test the user-template fill engine (pipeline/fill_template.py).

Generates a spread of realistic-and-messy user templates - correct markers and
headings, synonyms, typos, the WRONG {{ }} syntax, Chinese, inline markers,
duplicates, unrecognised fields - then runs each through the real engine
(filling with a chosen meeting) and reports what got recognised / placed /
skipped. The template + filled output for each go in tests/template_stress/ so
you can open both in Word and compare.

Usage:  python tests\\template_stress_test.py [meeting_id]   (default 28)
"""
from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

from docx import Document

BACKEND = Path(__file__).resolve().parent.parent
os.chdir(BACKEND)
sys.path.insert(0, str(BACKEND))
sys.stdout.reconfigure(encoding="utf-8")

from app.pipeline.fill_template import fill_user_document, recognised_fields  # noqa: E402

OUT = BACKEND / "tests" / "template_stress"
OUT.mkdir(parents=True, exist_ok=True)


def _details_table(doc, pairs):
    t = doc.add_table(rows=len(pairs), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(pairs):
        t.rows[i].cells[0].text = k
        t.rows[i].cells[1].text = v
    return t


# --- scenario builders ----------------------------------------------------

def d01_all_markers():
    """Every field as a [[marker]] - scalars inline in a table, blocks on their
    own line."""
    d = Document()
    d.add_paragraph("[ Acme Corp — logo here ]")
    d.add_heading("MEETING MINUTES", 0)
    d.add_paragraph("[[Meeting Title]]")
    _details_table(d, [("Date", "[[Date]]"), ("Duration", "[[Duration]]"),
                       ("Language", "[[Language]]"), ("Attendees", "[[Attendees]]")])
    for marker in ["[[Summary]]", "[[Key Decisions]]", "[[Action Items]]",
                   "[[Deadlines]]", "[[Technical Issues]]", "[[Risks]]", "[[Full Transcript]]"]:
        d.add_paragraph(marker)
    return d, "Pure [[markers]] for everything (scalars inline, blocks standalone)."


def d02_all_headings():
    """No markers at all - plain section headings the engine fills underneath."""
    d = Document()
    d.add_heading("Meeting Title", 1)
    for h in ["Date", "Duration", "Language", "Attendees", "Meeting Summary",
              "Key Decisions", "Action Items", "Deadlines", "Technical Issues",
              "Risks", "Full Transcript"]:
        d.add_heading(h, 1)
    return d, "Pure section HEADINGS, no markers."


def d03_mixed_typos_synonyms():
    """The messy real-world mix: a marker, valid synonyms, typos, the wrong
    {{ }} syntax, an unrecognised marker, a heading with a colon."""
    d = Document()
    d.add_paragraph("[[Meeting Title]]")            # marker (ok)
    d.add_heading("Overview", 1)                    # synonym -> summary (ok)
    d.add_heading("To-Do", 1)                       # synonym -> action_items (ok)
    d.add_heading("Resolutions", 1)                 # synonym -> decisions (ok)
    d.add_heading("Riks", 1)                        # TYPO of Risks -> NOT recognised
    d.add_heading("Acton Items", 1)                 # TYPO -> NOT recognised
    d.add_heading("Deadlines", 1)                   # exact (ok)
    d.add_heading("Blockers", 1)                    # synonym -> issues (ok)
    d.add_paragraph("{{ summary }}")                # WRONG syntax (docxtpl) -> left literal
    d.add_paragraph("[[Budget]]")                   # unrecognised marker -> left literal
    d.add_heading("Key Decisions:", 1)              # colon -> normalised; decisions already done -> skipped
    return d, "Mixed: marker + synonyms + TYPOS (Riks/Acton) + {{ }} mistake + [[Budget]] + colon heading."


def d04_chinese():
    """Chinese headings + markers."""
    d = Document()
    d.add_paragraph("[[会议标题]]")                  # title marker (zh)
    _details_table(d, [("日期", "[[日期]]"), ("时长", "[[Duration]]")])  # mixed zh/en marker
    for h in ["会议摘要", "决策", "行动项", "截止日期", "技术问题", "风险", "会议记录"]:
        d.add_heading(h, 1)
    return d, "Chinese headings (会议摘要/决策/行动项/风险…) + zh markers."


def d05_edge_cases():
    """Inline markers in a sentence, a marker in a table cell, duplicate
    heading, an embedded (non-exact) heading, sparse layout."""
    d = Document()
    d.add_paragraph("[[Meeting Title]]")
    d.add_paragraph("This session on [[Date]] ran for about [[Duration]] in [[Language]].")  # inline
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Attendees"
    t.rows[0].cells[1].text = "[[Attendees]]"        # marker inside a table cell
    d.add_heading("Action Items", 1)                 # fills
    d.add_heading("Action Items", 1)                 # DUPLICATE -> already filled -> skipped
    d.add_heading("Meeting Action Items List", 1)    # embedded words, NOT exact -> not recognised
    d.add_heading("Summary", 1)                      # fills
    return d, "Edge cases: inline markers, marker in a table cell, duplicate heading, non-exact heading."


def d06_wrong_syntax_only():
    """A user who copied docxtpl {{ }} syntax everywhere - the engine recognises
    NONE of it, so upload validation should REJECT this document."""
    d = Document()
    d.add_heading("MEETING MINUTES", 0)
    for line in ["{{ meeting.title }}", "{{ date }}", "{{ summary }}", "{{ action_items }}",
                 "{{ decisions }}", "{{ deadlines }}", "{{ issues }}", "{{ risks }}"]:
        d.add_paragraph(line)
    return d, "ONLY {{ }} docxtpl syntax (wrong for user templates) -> 0 recognised -> should be REJECTED."


SCENARIOS = [
    ("01_all_markers", d01_all_markers),
    ("02_all_headings", d02_all_headings),
    ("03_mixed_typos_synonyms", d03_mixed_typos_synonyms),
    ("04_chinese", d04_chinese),
    ("05_edge_cases", d05_edge_cases),
    ("06_wrong_syntax_only", d06_wrong_syntax_only),
]


def main() -> int:
    meeting_id = int(sys.argv[1]) if len(sys.argv) > 1 else 28
    print(f"=== Template stress test (filling with meeting {meeting_id}) ===")
    print(f"templates + filled outputs -> {OUT}\n")
    for name, build in SCENARIOS:
        doc, desc = build()
        tmpl = OUT / f"{name}_TEMPLATE.docx"
        doc.save(str(tmpl))
        print(f"--- {name} ---")
        print(f"    {desc}")
        recog = recognised_fields(tmpl)
        print(f"    recognised fields: {recog or '(none)'}")
        if not recog:
            print("    => REJECTED at upload (no recognised heading/marker). [validation works]\n")
            continue
        out_path, report = fill_user_document(meeting_id, tmpl)
        filled = OUT / f"{name}_FILLED.docx"
        shutil.copy(out_path, filled)
        placed = ", ".join(f"{k}({v})" for k, v in report["placed"].items()) or "(none)"
        skipped = ", ".join(report["skipped"]) or "(none)"
        print(f"    placed:  {placed}")
        print(f"    skipped (had data, no spot): {skipped}")
        print(f"    -> {filled.name}\n")
    print("Done. Open the *_TEMPLATE.docx (input) and *_FILLED.docx (output) pairs to compare.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
